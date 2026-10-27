"""
Build the BCSE497J project report on top of the college template.

    python docs/report/build_report.py            # builds, paginates with Word, rebuilds with real page numbers, exports PDF

Run with the default workspace (do not set ATERNET_WORKSPACE). Needs Microsoft Word for pagination;
without it the document is still produced, but page numbers in the contents and lists stay at 0 until
the fields are updated in Word (Ctrl+A, F9).
"""

import copy
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

from report_data import FIG, REFS, ROOT, d
from report_lib import (C_, J, L_, R_, Report, roman, set_font)
import report_text1 as T1
import report_text2 as T2
import report_text3 as T3

TEMPLATE = ROOT / "BCSE497J Project I Report - Template FINAlL Updated  (2).docx"
OUT = ROOT / "ATERNET_Project_Report_BCSE497J.docx"
PDF = ROOT / "ATERNET_Project_Report_BCSE497J.pdf"

TITLE = T1.TITLE.upper()
PROGRAMME = "Computer Science and Engineering"
# Fields only the student can supply; kept highlighted in yellow so they cannot be missed.
PH_TEAM = [("<Reg. No.>", "<STUDENT NAME 1>"), ("<Reg. No.>", "<STUDENT NAME 2>"), ("<Reg. No.>", "<STUDENT NAME 3>")]
PH_GUIDE, PH_DESIG = "<Project Guide Name>", "<Designation>"
PH_HOD, PH_DEPT = "<Head of the Department's Name>", "<Department Name>"
PH_STUDENTS = "<Student Name(s) (Reg. No.)>"
PH_CANDIDATE = "<Name of the Candidate(s)>"


# ------------------------------------------------------------------ small helpers
def paragraphs(doc):
    return list(doc.paragraphs)


def find(doc, text, exact=False, nth=0):
    hits = [p for p in doc.paragraphs if (p.text.strip() == text if exact else text in p.text)]
    return hits[nth] if len(hits) > nth else None


def delete(p):
    p._element.getparent().remove(p._element)


def clear(p):
    for ch in list(p._p):
        if ch.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:fldSimple"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p._p.remove(ch)


def rebuild(p, segs, size=12, align=None, line=None, after=None):
    """Replace a paragraph's content. segs: list of str or (str, flags) with flags in 'b','i','h' (highlight)."""
    clear(p)
    for seg in segs:
        text, flags = (seg, "") if isinstance(seg, str) else seg
        r = p.add_run(text)
        set_font(r, size, "b" in flags or None, "i" in flags or None, highlight="h" in flags)
    if align is not None:
        p.alignment = align
    if line:
        p.paragraph_format.line_spacing = line
    if after is not None:
        p.paragraph_format.space_after = Pt(after)


def set_cell(cell, text, size=12, bold=False, align=None, left=None, highlight=False):
    for extra in cell.paragraphs[1:]:
        delete(extra)
    p = cell.paragraphs[0]
    clear(p)
    r = p.add_run(text)
    set_font(r, size, bold, highlight=highlight)
    if align is not None:
        p.alignment = align
    p.paragraph_format.left_indent = Cm(left) if left else None
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
    return p


def clone_row(table, src_idx):
    new = copy.deepcopy(table.rows[src_idx]._tr)
    table._tbl.append(new)
    return table.rows[-1]


def drop_rows(table, keep):
    for r in list(table.rows)[keep:]:
        table._tbl.remove(r._tr)


# ------------------------------------------------------------------ front matter
def compact_title_page(doc):
    """Remove the template's empty spacer paragraphs on the title page so title, team, guide, logo and date fit one page."""
    for el in list(doc.element.body.iterchildren()):
        if el.tag != qn("w:p"):
            continue
        ppr = el.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            break                                    # end of the title-page section
        text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
        has_drawing = bool(el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:pict")))
        if not text and not has_drawing:
            el.getparent().remove(el)
    spacing = {"Submitted in partial": 20, "Bachelor of Technology": 8, "Under the Supervision": 10}
    for p in doc.paragraphs[:20]:
        t = p.text.strip()
        for key, pts in spacing.items():
            if t.startswith(key):
                p.paragraph_format.space_before = Pt(pts)
        if t == "by":
            p.paragraph_format.space_before = Pt(8)
        if p._p.findall(".//" + qn("w:drawing")) or p._p.findall(".//" + qn("w:pict")):
            p.paragraph_format.space_before = Pt(14)
            p.alignment = C_
    doc.paragraphs[0].paragraph_format.space_after = Pt(10)


def front_matter(doc, R):
    compact_title_page(doc)
    # --- title page
    p = find(doc, "<Title of the Project>")
    rebuild(p, [(TITLE, "b")], size=18, align=C_, line=1.5)
    for t in ("(Times New Roman 18", "(mention names of team members"):
        q = find(doc, t)
        if q is not None:
            delete(q)
    q = find(doc, "Programme", exact=True)
    rebuild(q, [(PROGRAMME, "b")], size=19, align=C_)
    t0 = doc.tables[0]
    for row, (reg, nm) in zip(t0.rows, PH_TEAM):
        set_cell(row.cells[0], reg, 13, True, highlight=True)
        set_cell(row.cells[1], nm, 13, True, highlight=True)
    t1 = doc.tables[1]
    set_cell(t1.rows[0].cells[0], PH_GUIDE, 14, True, align=C_, highlight=True)
    set_cell(t1.rows[1].cells[0], PH_DESIG, 14, False, align=C_, highlight=True)
    set_cell(t1.rows[2].cells[0], "School of Computer Science and Engineering (SCOPE)", 14, False, align=C_)

    # --- declaration
    p = find(doc, "hereby")
    rebuild(p, ["I hereby declare that the project entitled ", (T1.TITLE, "b"),
                " submitted by me, for the award of the degree of ", ("Bachelor of Technology in Computer Science and Engineering", "i"),
                " to VIT is a record of bonafide work carried out by me under the supervision of Prof. / Dr. ",
                (PH_GUIDE, "h"), "."], align=J, line=1.5)
    q = find(doc, "I further declare")
    rebuild(q, ["I further declare that the work reported in this project has not been submitted and will not be "
                "submitted, either in part or in full, for the award of any other degree or diploma in this institute "
                "or any other institute or university."], align=J, line=1.5)
    for q in [x for x in doc.paragraphs if x.text.startswith("Note:")]:
        delete(q)

    # --- certificate
    p = find(doc, "This is to certify")
    rebuild(p, ["This is to certify that the project entitled ", (T1.TITLE, "b"),
                " submitted by ", (PH_STUDENTS, "bh"),
                ", School of Computer Science and Engineering, VIT, for the award of the degree of ",
                ("Bachelor of Technology in Computer Science and Engineering", "i"),
                ", is a record of bonafide work carried out by the candidate(s) under my supervision during Fall "
                "Semester 2026-2027, as per the VIT code of academic and research ethics."], align=J, line=1.5)
    q = find(doc, "The contents of this report have not been submitted")
    rebuild(q, ["The contents of this report have not been submitted and will not be submitted either in part or in "
                "full, for the award of any other degree or diploma in this institute or any other institute or "
                "university. The project fulfills the requirements and regulations of the University and in my opinion "
                "meets the necessary standards for submission."], align=J, line=1.5)
    q = find(doc, "<< Head of the Department >>")
    rebuild(q, [(PH_HOD, "bh")], align=C_, line=1.5)
    q = find(doc, "<< Programme >>")
    rebuild(q, [(PROGRAMME, "b")], align=C_, line=1.5)

    # --- acknowledgements
    p = find(doc, "I express my profound appreciation")
    rebuild(p, ["I express my profound appreciation to ", (PH_HOD, "h"), ", the Head of the ", (PH_DEPT, "h"),
                ", for his/her insightful guidance and continuous support. His/her expertise and advice have been crucial in "
                "shaping the direction of my project. The Head of Department’s commitment to fostering a collaborative and "
                "supportive atmosphere has greatly enhanced my learning experience. His/her constructive feedback and "
                "encouragement have been invaluable in overcoming challenges and achieving my project goals."],
            align=J, line=1.5)
    p = find(doc, "I am immensely thankful")
    rebuild(p, ["I am immensely thankful to my project guide, ", (PH_GUIDE, "h"),
                ", for his/her dedicated mentorship and invaluable feedback. His/her patience, knowledge, and encouragement "
                "have been pivotal in the successful completion of this project. My supervisor’s willingness to share his/her "
                "expertise and provide thoughtful guidance has been instrumental in refining my ideas and methodologies. "
                "His/her support has not only contributed to the success of this project but has also enriched my overall "
                "academic experience."], align=J, line=1.5)
    q = find(doc, "Name of the Candidate", exact=True)
    rebuild(q, [(PH_CANDIDATE, "bh")])

    # --- table of contents: drop the instruction line, rebuild the table
    q = find(doc, "<Contents, Times New Roman 12, Line spacing 1.5>")
    if q is not None:
        delete(q)
    fill_toc(doc.tables[2], R)

    # --- list of figures
    q = find(doc, "(In the chapters, figure caption should come")
    if q is not None:
        delete(q)
    fill_list(doc.tables[3], R, R.figs, "Figure No.")

    # --- list of tables: replace the tab-separated sample lines by a real table
    h = find(doc, "List of Tables", exact=True)
    for t in ("Table No.", "Table caption"):
        q = find(doc, t)
        if q is not None:
            delete(q)
    new_tbl = copy.deepcopy(doc.tables[3]._tbl)
    h._p.addnext(new_tbl)
    from docx.table import Table
    lot = Table(new_tbl, h._parent)
    fill_list(lot, R, R.tabs, "Table No.")
    spacer = h._p.getnext()
    # blank paragraph between heading and table
    # (keep the template's own empty paragraph that followed the heading)

    # --- abbreviations (alphabetical)
    abbr = sorted(ABBREVIATIONS.items(), key=lambda kv: kv[0].lower())
    first = find(doc, "2G\t")
    proto = copy.deepcopy(first._p)
    for q in [x for x in doc.paragraphs if re.match(r"^(2G|3GPP|3G|4G|AWGN|BBE|CSMA|DAS)\t", x.text)]:
        delete(q)
    # insert after the 'List of Abbreviations' heading's blank paragraph
    head = find(doc, "List of Abbreviations", exact=True)
    anchor = head._p.getnext()
    for k, (a, full) in enumerate(abbr):
        el = copy.deepcopy(proto)
        anchor.addnext(el)
        anchor = el
        from docx.text.paragraph import Paragraph
        pp = Paragraph(el, head._parent)
        rebuild(pp, [a + "\t" + full], size=12)
        hang(pp, 3.6)
    # remove the empty spacer paragraphs and the Note that followed the sample list
    # (Note paragraphs were removed above; blank ones between abbreviations and the page break are harmless)

    # --- symbols and notations
    sym_head = find(doc, "Symbols and Notations", exact=True)
    abs_head = find(doc, "ABSTRACT", exact=True)
    el = sym_head._p.getnext()
    while el is not None and el is not abs_head._p:
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt
    anchor = sym_head._p
    for s_, desc in SYMBOLS:
        e2 = copy.deepcopy(proto)
        anchor.addnext(e2)
        anchor = e2
        pp = Paragraph(e2, sym_head._parent)
        rebuild(pp, [s_ + "\t" + desc], size=12)
        hang(pp, 6.4)
    # page break before the abstract
    abs_head.paragraph_format.page_break_before = True

    # --- abstract
    rebuild(abs_head, [("ABSTRACT", "b")], size=16, align=C_)
    R._bookmark(abs_head, "h_abstract")
    el = abs_head._p.getnext()
    break_p = None
    while el is not None:
        nxt = el.getnext()
        ppr = el.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            break_p = el
            break
        el.getparent().remove(el)
        el = nxt
    anchor = abs_head._p
    words = 0
    for text in T1.abstract_paragraphs():
        e2 = OxmlElement("w:p")
        anchor.addnext(e2)
        anchor = e2
        pp = Paragraph(e2, abs_head._parent)
        pp.alignment = J
        pp.paragraph_format.line_spacing = 1.15
        pp.paragraph_format.space_after = Pt(6)
        pp.paragraph_format.space_before = Pt(0) if pp is not None else None
        r = pp.add_run(text)
        set_font(r, 12)
        words += len(text.split())
    print(f"abstract: {words} words")
    assert 300 <= words <= 500, f"abstract has {words} words (template: 300-500)"
    return break_p


ABBREVIATIONS = {
    "ACE": "Average Coverage Error",
    "API": "Application Programming Interface",
    "ATERNET": "Adaptive Temporal Expert Routing Network",
    "CQR": "Conformalized Quantile Regression",
    "CRPS": "Continuous Ranked Probability Score",
    "DFD": "Data Flow Diagram",
    "EMA": "Exponential Moving Average",
    "ERA5": "ECMWF Reanalysis, fifth generation",
    "GBDT": "Gradient-Boosted Decision Trees",
    "GEM": "Global Environmental Multiscale model",
    "GFS": "Global Forecast System",
    "GHI": "Global Horizontal Irradiance",
    "ICON": "Icosahedral Nonhydrostatic model",
    "LSTM": "Long Short-Term Memory",
    "MAE": "Mean Absolute Error",
    "MLP": "Multi-Layer Perceptron",
    "NWP": "Numerical Weather Prediction",
    "PICP": "Prediction Interval Coverage Probability",
    "PINAW": "Prediction Interval Normalised Average Width",
    "PV": "Photovoltaic",
    "R²": "Coefficient of determination",
    "RMSE": "Root Mean Square Error",
    "UKMO": "United Kingdom Meteorological Office model",
    "nMAE": "Normalised Mean Absolute Error",
    "nRMSE": "Normalised Root Mean Square Error",
}
SYMBOLS = [
    ("α", "Mixing weight between the two uncertainty signals (0 to 1)"),
    ("β", "Learned sensitivity of the gate temperature to the uncertainty"),
    ("γ", "Learned sensitivity of the interval width to the uncertainty"),
    ("δ", "Learned scale of the residual gate (initialised to 0)"),
    ("σ_NWP", "Spread between the weather-forecast models"),
    ("σ_expert", "Disagreement between the four experts"),
    ("σ_total", "Combined uncertainty used by the gate and the decoder"),
    ("T", "Gate temperature"),
    ("τ", "Quantile level of the pinball loss (0.1, 0.5, 0.9)"),
    ("ŷ", "Forecast of an expert or of the router"),
    ("h_expert, h_sigma, h_nwp", "Embeddings produced by the backbone"),
    ("P10, P50, P90", "10th, 50th and 90th percentile forecasts"),
    ("X_seq, X_nwp, X_sigma", "Input tensors (history, forecast weather, forecast-spread features)"),
    ("θ_ema", "Exponentially averaged model weights"),
]


def hang(p, cm):
    """Term in the first column, description in the second; wrapped lines align under the description."""
    pf = p.paragraph_format
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:tabs")):
        ppr.remove(old)
    pf.left_indent = Cm(cm)
    pf.first_line_indent = Cm(-cm)
    pf.tab_stops.add_tab_stop(Cm(cm))
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0


def fill_toc(table, R):
    """Rebuild the Sl.No | Contents | Page No. table from the headings collected while writing the chapters."""
    hdr = table.rows[0]
    chapter_tpl = copy.deepcopy(table.rows[2]._tr)
    section_tpl = copy.deepcopy(table.rows[3]._tr)
    drop_rows(table, 1)

    def add(tpl, sl, text, bm, bold, left=None):
        table._tbl.append(copy.deepcopy(tpl))
        row = table.rows[-1]
        set_cell(row.cells[0], sl, 12, bold)
        set_cell(row.cells[1], text, 12, bold, left=left)
        cell = row.cells[2]
        p = cell.paragraphs[0]
        clear(p)
        p.alignment = C_
        R.pageref(p, bm, 12, bold)

    add(chapter_tpl, "", "Abstract", "h_abstract", True)
    for level, sl, text, bm in R.toc:
        if level == 0:
            add(chapter_tpl, sl, text.upper() if sl else text, bm, True)
        elif level == 1:
            add(section_tpl, "", text, bm, False)
        else:
            add(section_tpl, "", text, bm, False, left=0.8)


def fill_list(table, R, items, first_col):
    header_tr = copy.deepcopy(table.rows[0]._tr)
    row_tpl = copy.deepcopy(table.rows[1]._tr)
    set_cell(table.rows[0].cells[0], first_col, 14, True)
    drop_rows(table, 1)
    for label, caption, bm in items:
        table._tbl.append(copy.deepcopy(row_tpl))
        row = table.rows[-1]
        set_cell(row.cells[0], label, 12)
        set_cell(row.cells[1], re.sub(r"\[@[^\]]+\]", "", caption), 12)
        p = row.cells[2].paragraphs[0]
        clear(p)
        p.alignment = C_
        R.pageref(p, bm, 12)


# ------------------------------------------------------------------ main matter
def main_matter(doc, R, break_p):
    # remove everything after the end-of-front-matter section break (the template's sample chapter text)
    el = break_p.getnext()
    while el is not None:
        nxt = el.getnext()
        if el.tag == qn("w:sectPr"):
            break
        el.getparent().remove(el)
        el = nxt
    # numbering of the main matter restarts at 1
    sect = doc.sections[-1]._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sect.append(pg)

    shots = make_screenshot_crops()
    T3_ = T3
    T1.chapter1(R)
    T1.chapter2(R)
    T1.chapter3(R)
    T2.chapter4(R)
    T2.chapter5(R)
    T3.chapter6(R, shots)
    T3.chapter7(R)
    T3.chapter8(R)
    references(R)
    T3.appendix(R)


def _sidebar_x(im):
    """x position where the dark sidebar ends (the sidebar colour is #2E3440)."""
    import numpy as np
    row = np.asarray(im.convert("RGB"))[600]
    for x in range(100, 600):
        if int(row[x].sum()) > 450:
            return x
    return 272


def _split_chunks(im, x0, max_h=1180, min_h=450):  # a merged last piece may reach ~1500 px
    """Split a tall screenshot at blank (white) rows so no line of text is cut."""
    import numpy as np
    gray = np.asarray(im.convert("L"))[:, x0:]
    chunks, top = [], 0
    while im.height - top > max_h:
        cut = None
        for y in range(top + max_h, top + min_h, -1):
            if (gray[y - 8:y + 8] >= 246).all():
                cut = y
                break
        cut = cut or top + max_h
        chunks.append((top, cut))
        top = cut
    chunks.append((top, im.height))
    if len(chunks) > 1 and chunks[-1][1] - chunks[-1][0] < 320:      # no near-empty last piece
        chunks[-2] = (chunks[-2][0], chunks[-1][1])
        chunks.pop()
    return chunks


def make_screenshot_crops():
    """Cut the dashboard screenshots into page-sized pieces; the overview's first piece keeps the menu."""
    out = {}
    pages = {
        "overview": "shot_6_1_overview.png", "better": "shot_6_2_is_aternet_better.png",
        "india": "shot_6_5_india_demo.png", "sure": "shot_6_3_how_sure.png",
        "safety": "shot_6_4_safety_checks.png", "explore": "shot_6_6_explore.png",
        "experts": "shot_6_7_experts.png",
    }
    for key, fname in pages.items():
        im = Image.open(FIG / fname).convert("RGB")
        x0 = _sidebar_x(im)
        paths = []
        for k, (t, b) in enumerate(_split_chunks(im, x0)):
            keep_menu = (key == "overview" and k == 0)
            piece = im.crop((0 if keep_menu else x0, max(t, 30 if k == 0 else t), im.width, b))
            dst = FIG / f"crop_{key}_{k + 1}.png"
            piece.save(dst)
            paths.append(dst)
        out[key] = paths
    return out


def references(R):
    R.chapter_heading(9, "References")
    R.body("References are numbered in the order in which they are first cited (IEEE style).", after=8)
    for key, n in sorted(R.cites.items(), key=lambda kv: kv[1]):
        p = R.para("", size=12, after=6, spacing=1.15, left_indent=1.0, first_indent=-1.0)
        r = p.add_run(f"[{n}]\t"); set_font(r, 12)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
        R.add_text(p, REFS[key].replace("*", ""), 12)


# ------------------------------------------------------------------ build + paginate
def build(pages):
    doc = Document(str(TEMPLATE))
    R = Report(doc, pages)
    R.ref_text = REFS
    R.chapter = 0
    # chapter bodies first (they fill R.toc / R.figs / R.tabs), but they must come after the template's
    # front matter in the file; front-matter edits only need the lists, so run main matter first.
    break_p = None
    # find the end-of-front-matter break paragraph before editing anything
    marks = [p._p for p in doc.paragraphs if p._p.find(qn("w:pPr")) is not None
             and p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
    break_p = marks[1]          # ends the roman-numbered front matter
    main_matter(doc, R, break_p)
    front_matter(doc, R)
    global OUT, PDF
    try:
        doc.save(str(OUT))
    except PermissionError:
        # The report is open in Word: keep the open file untouched and write a copy next to it.
        OUT = OUT.with_name(OUT.stem.replace("_updated", "") + "_updated.docx")
        PDF = OUT.with_suffix(".pdf")
        doc.save(str(OUT))
        print(f"NOTE: the original is open in Word, so the new version was saved as {OUT.name}")
    return R


def paginate(R, export_pdf=False):
    import win32com.client
    w = win32com.client.DispatchEx("Word.Application")
    w.Visible = False
    w.DisplayAlerts = 0
    try:
        doc = w.Documents.Open(str(OUT), ReadOnly=False)
        doc.Repaginate()
        pages = {}
        for bm in R.bookmarks:
            if not doc.Bookmarks.Exists(bm):
                continue
            rng = doc.Bookmarks(bm).Range
            n = int(rng.Information(1))
            sec = rng.Sections(1).Index
            pages[bm] = roman(n) if sec == 2 else str(n)
        total = int(doc.ComputeStatistics(2))
        if export_pdf:
            try:
                doc.ExportAsFixedFormat(str(PDF), 17)
            except Exception:
                alt = PDF.with_name(PDF.stem + "_new.pdf")
                doc.ExportAsFixedFormat(str(alt), 17)
                print(f"NOTE: {PDF.name} is in use; PDF saved as {alt.name}")
        doc.Close(False)
    finally:
        w.Quit()
    return pages, total


if __name__ == "__main__":
    R = build({})
    pages, total = paginate(R)
    print(f"pass 1: {total} pages, {len(pages)} bookmarks")
    R = build(pages)
    pages2, total2 = paginate(R, export_pdf=True)
    changed = {k: (pages[k], pages2[k]) for k in pages2 if pages.get(k) != pages2[k]}
    print(f"pass 2: {total2} pages; page numbers that moved between passes: {len(changed)}")
    if changed:      # one more round settles any repagination caused by wider/narrower numbers
        R = build(pages2)
        pages3, total3 = paginate(R, export_pdf=True)
        print(f"pass 3: {total3} pages; still moving: {sum(1 for k in pages3 if pages2.get(k) != pages3[k])}")
    print("figures:", [f[0] for f in R.figs])
    print("tables :", [t[0] for t in R.tabs])
    print("citations:", len(R.cites))
