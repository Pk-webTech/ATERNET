"""
Build the Review-3 presentation on top of the college template.

    python docs/ppt/build_review_ppt.py

Input : "Review 3- 28-10- 2026.pptx" (the template, not modified)
Output: "Review3_ATERNET_28-10-2026.pptx"   (19 slides)

Numbers come from docs/results_summary.json; figures from docs/report/figures.
"""

import copy
import json
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / "docs" / "report" / "figures"
TEMPLATE = ROOT / "Review 3- 28-10- 2026.pptx"
OUT = ROOT / "Review3_ATERNET_28-10-2026.pptx"
FONT = "Times New Roman"
SUMMARY = json.loads((ROOT / "docs" / "results_summary.json").read_text(encoding="utf-8"))

TITLE = "ATERNET: Adaptive Temporal Expert Routing Network for Probabilistic Solar Power Forecasting"
BLUE = RGBColor(0x2E, 0x5A, 0x88)
RED = RGBColor(0xB8, 0x32, 0x32)


# ------------------------------------------------------------------ text helpers
def _highlight(run):
    rpr = run._r.get_or_add_rPr()
    hl = etree.SubElement(rpr, qn("a:highlight"))
    clr = etree.SubElement(hl, qn("a:srgbClr"))
    clr.set("val", "FFFF00")
    rpr.remove(hl)
    # schema order: ... fill, effect, highlight, ..., latin, ea, cs, sym ...
    anchor = None
    for tag in ("a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick"):
        anchor = rpr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(hl)
    else:
        rpr.append(hl)


def style_run(run, size=None, bold=None, italic=None, color=None, highlight=False, baseline=None):
    f = run.font
    f.name = FONT
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if baseline:
        run._r.get_or_add_rPr().set("baseline", str(baseline))
    if highlight:
        _highlight(run)


_TOKEN = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\}|_[A-Za-z0-9]|\^[A-Za-z0-9+\-]|[^_^]+)")


def add_rich(p, text, size, bold=False, italic=False, color=None, highlight=False):
    """Text with _{sub} / ^{sup} markers (single character: _k, ^T)."""
    for tok in _TOKEN.findall(text):
        if tok[0] == "_":
            r = p.add_run()
            r.text = tok[1:].strip("{}")
            style_run(r, size * 0.9, bold, italic, color, highlight, baseline=-25000)
        elif tok[0] == "^":
            r = p.add_run()
            r.text = tok[1:].strip("{}")
            style_run(r, size * 0.9, bold, italic, color, highlight, baseline=30000)
        else:
            r = p.add_run()
            r.text = tok
            style_run(r, size, bold, italic, color, highlight)


def set_paragraphs(tf, items, size=18, level_sizes=None, bullets=True, space_before=6):
    """items: list of str | (str, level) | (str, level, dict(bold=..., highlight=...))."""
    tf.clear()
    first = True
    for it in items:
        text, level, opt = (it, 0, {}) if isinstance(it, str) else (it[0], it[1], it[2] if len(it) > 2 else {})
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_before = Pt(opt.get("before", space_before))
        sz = (level_sizes or {}).get(level, size)
        add_rich(p, text, opt.get("size", sz), opt.get("bold", False), opt.get("italic", False), opt.get("color"), opt.get("highlight", False))
        if not bullets or opt.get("nobullet"):
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "0")
            ppr.set("indent", "0")
            for old in ppr.findall(qn("a:buNone")):
                ppr.remove(old)
            etree.SubElement(ppr, qn("a:buNone"))
    return tf


def remove(shape):
    shape._element.getparent().remove(shape._element)


# ------------------------------------------------------------------ slide helpers
def footers_from(src_slide):
    return [copy.deepcopy(sh._element) for sh in src_slide.placeholders
            if sh.placeholder_format.type in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)]


def new_slide(prs, title, foot_elems, body=False):
    layout = prs.slide_layouts[1]
    s = prs.slides.add_slide(layout)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Cm(1.3), Cm(1.9), Cm(29.2), Cm(1.6)
    t.text_frame.clear()
    r = t.text_frame.paragraphs[0].add_run()
    r.text = title
    style_run(r, 30, True)
    if not body:
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx == 1:
                remove(ph)
    for el in foot_elems:
        s.shapes._spTree.append(copy.deepcopy(el))
    return s


def set_title(slide, text, size=30):
    t = slide.shapes.title
    t.text_frame.clear()
    r = t.text_frame.paragraphs[0].add_run()
    r.text = text
    style_run(r, size, True)


def body_ph(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            return ph
    return None


def place(ph, left, top, width, height):
    ph.left, ph.top, ph.width, ph.height = Cm(left), Cm(top), Cm(width), Cm(height)


def add_picture_fit(slide, path, left, top, max_w, max_h, center=True):
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = w * scale, h * scale
    x = left + (max_w - pw) / 2 if center else left
    y = top + (max_h - ph) / 2 if center else top
    return slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(pw), Cm(ph))


def add_text(slide, left, top, width, height, items, size=14, bullets=False, anchor=MSO_ANCHOR.TOP, align=None):
    tb = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    set_paragraphs(tf, items, size=size, bullets=bullets, space_before=3)
    if align is not None:
        for p in tf.paragraphs:
            p.alignment = align
    return tb


def add_table(slide, left, top, col_w, rows, size=10, row_h=None, rich_cols=(), bold_first_col=False, header_size=None):
    n_r, n_c = len(rows), len(col_w)
    gf = slide.shapes.add_table(n_r, n_c, Cm(left), Cm(top), Cm(sum(col_w)), Cm(row_h * n_r if row_h else n_r * 1.0))
    tbl = gf.table
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Cm(w)
    for ri, row in enumerate(rows):
        if row_h:
            tbl.rows[ri].height = Cm(row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Cm(0.12)
            cell.margin_top = cell.margin_bottom = Cm(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.clear()
            lines = val.split("\n") if isinstance(val, str) else [str(val)]
            for li, line in enumerate(lines):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                bold = ri == 0 or (bold_first_col and ci == 0)
                sz = (header_size or size) if ri == 0 else size
                if ci in rich_cols and ri > 0:
                    add_rich(p, line, sz, bold=bold, italic=False)
                else:
                    r = p.add_run()
                    r.text = line
                    style_run(r, sz, bold)
    return gf


# ------------------------------------------------------------------ data
def rows(site):
    return {r["model"].strip(): r for r in SUMMARY["sites"][site]["rows"]}


Y, SW = rows("yorkshire"), rows("south_west")
A = "ATERNET (final: residual gate + shortcut + EMA)"


def ms(r, key, nd=2):
    return f"{r[key]['mean']:.{nd}f} ± {r[key]['std']:.{nd}f}"


# ------------------------------------------------------------------ literature rows (from the PDFs' abstracts)
LIT = [
    ["Paper", "Objective", "Methodology", "Pros / Cons (for our goal)", "Findings (as reported)"],
    ["[1] Wang et al., 2026\nIEEE Trans. Ind. Appl.", "Distributed PV, ultra-short-term, deterministic + probabilistic",
     "Dynamic graph network; shape-amplitude loss (DTW + MSE); quantile regression",
     "+ models site correlation and shape\n− needs many sites", "Validated on real distributed-PV data; shape and amplitude kept consistent"],
    ["[2] Lo et al., 2025\nIEEE Trans. Ind. Appl.", "Faster, simpler probabilistic solar forecasts",
     "Quantile regression; skips training bounds when output is stable; imputation",
     "+ cheaper to run\n− no uncertainty source modelled", "Less forecasting time and better accuracy than compared methods"],
    ["[3] He et al., 2026\nICSGGE (IEEE)", "PV intervals under complex weather",
     "Adaptive VMD + BiLSTM + feature attention; monotonicity penalty",
     "+ handles non-stationarity\n− penalty, not guarantee, against crossing", "Better interval reliability in cloudy / convective weather; 141.6 ms run time"],
    ["[4] Wang et al., 2025\nICEMCE (IEEE)", "Robust, explainable PV ensemble",
     "Context-gated mixture of experts; residual corrector; distillation; conformal + MC-dropout",
     "+ closest gating idea\n− no NWP, no NWP-spread signal", "Outperforms baselines and ensembles on a multi-location PV dataset"],
    ["[5] Zhao et al., 2026\nIEEE Trans. Power Syst.", "Use several NWP sources for probabilistic PV",
     "Two-stage ensemble, meta-learner, feature + temporal attention, non-crossing quantiles",
     "+ multi-source NWP as input\n− spread between NWPs not used as uncertainty", "Better reliability, sharpness and skill; multi-source NWP lowers single-source risk"],
    ["[6] He et al., 2026\nIEEE Trans. Ind. Informat.", "Day-ahead PV probability density",
     "Similar-moment selection; GAN on measurements + NWP; CNN-LSTM; QR + KDE",
     "+ enriches training data\n− GAN training complexity", "Lowest CRPS in most weather scenarios; +45.8% point / +51.5% probabilistic vs QRTransformer"],
    ["[7] Heidrich et al., 2026\nIEEE Access (curtailment)", "Curtailment forecast with calibrated intervals",
     "Patch Transformer; SwiGLU mixture of experts + load-balancing loss; conformal prediction",
     "+ same building blocks as ours\n− different series (not PV)", "MSE 0.0059 at 3 h (−13.2%); 91.27% coverage at nominal 90%"],
    ["[8] Lyu & Eftekharnejad, 2024\nIEEE Access", "Probabilistic solar under rapidly changing weather",
     "Copula theory + XGBoost weather classifier; dynamic spatio-temporal correlations",
     "+ weather-dependent modelling\n− heavy multi-variable setup", "Up to 60% higher accuracy in non-sunny autumn / winter"],
    ["[9] Lin et al., 2025\nIEEE Trans. Power Syst.", "Regional PV from many plants",
     "Dynamic convolutional graph network, graph pooling, parameter transfer (127 plants)",
     "+ end-to-end regional forecast\n− needs plant-level data", "High-quality, reliable regional probabilistic forecasts on public data"],
    ["[10] Huang et al., 2026\nIEEE Trans. Ind. Appl.", "Distributed-PV cluster, next 4 hours",
     "Spatial-temporal attention + graph convolution; non-crossing joint quantile loss",
     "+ multiple levels at once\n− graph over sub-regions required", "Verified on real-world datasets"],
]
LIT2 = [
    LIT[0],
    ["[11] Rastgoo et al., 2025\nIEEE Access", "Ultra-short-term PV from sky images",
     "Diffusion model; cross-attention on multi-resolution image patches; ELBO loss",
     "+ sees cloud motion\n− needs sky-image stream", "Superior to 64 baselines on four metrics"],
    ["[12] Seman et al., 2026\nIEEE Access (thermal energy)", "Thermal-energy forecasting",
     "Hierarchical VMD, architecture search, sparse expert Transformer with load balancing",
     "+ expert routing at scale\n− different series", "MAPE 5.94%, about 12–15% lower than competing methods"],
    ["[13] Phan et al., 2024\nIEEE Trans. Ind. Appl.", "One-day-ahead PV intervals",
     "Transformer-LUBE; XGBoost imputation; GRU residual correction; NWP",
     "+ missing-data handling\n− bounds, not quantiles", "Better NRMSE, NMAPE, PICP, PINAW, CWC than ANN, LSTM, GRU, XGBoost"],
    ["[14] Jensen et al., 2024\nIEEE TNNLS", "Valid intervals for any time-series model",
     "Ensemble conformalized quantile regression; asymmetric conformity scores",
     "+ distribution-free validity\n− generic, not PV-specific", "Sharper, valid intervals than QR-only or conformal-only"],
    ["[15] Song et al., 2025\nIEEE Trans. Sustain. Energy", "Large-scale multi-site PV with missing data",
     "Random coarse-graph attention; adaptive imputation (1,649 sites)",
     "+ scalable\n− multi-site data needed", "7–10% and 6–25% average improvement under missing-data types"],
    ["[16] Wang et al., 2026\nIEEE Trans. Sustain. Energy", "PV from satellite cloud forecasts",
     "Diffusion image prediction + neural SDE; multimodal ensemble with PV + NWP",
     "+ models cloud uncertainty\n− needs satellite images", "180-min CRPS −6.06% and −4.73% at two sites; 36.5% lower CRPS in images"],
    ["[17] Zhao et al., 2024\nIEEE Trans. Ind. Informat. (load)", "Probabilistic multi-energy load",
     "Hybrid-attention Transformer; Gaussian-process residual learning",
     "+ residual modelling\n− load, not PV", "Captures coupling between energy loads; probabilistic output"],
    ["[18] Ye et al., 2024\nIEEE Trans. Power Syst.", "Multi-horizon probabilistic energy",
     "TADNet: regular + residual components; convolution + attention; gating",
     "+ interval from residual\n− energy series generally", "Best performance on two energy datasets; robust long-term"],
    ["[19] Wang et al., 2026\nIEEE Trans. Ind. Appl.", "Distributed-PV intervals from joint distribution",
     "Transformer point forecast; autoencoder fluctuation patterns; Clayton copula",
     "+ pattern-dependent intervals\n− two-step, copula fitting", "Intervals from point forecast and error joint distribution"],
    ["[20] Abedinia et al., 2024\nIEEE Trans. Ind. Appl. (wind)", "Probabilistic wind forecast",
     "CNN-LSTM predicts spline quantile-function parameters; adaptive weighted CRPS",
     "+ full distribution\n− wind, not PV", "More accurate than compared models across scenarios"],
]

# ------------------------------------------------------------------ formulas (position -> formula)
FORMULAS = [
    # (id, position in the system, formula (rich text), what it does)
    ("F1", "Data preparation (Phase 1)",
     "pv_{norm} = P / P_{capacity}\nCSI = GHI / GHI_{clear-sky}\nx′ = (x − x_{min}) / (x_{max} − x_{min})",
     "Capacity-normalised target, clear-sky index feature, min-max scaling (fitted on train only)"),
    ("F2", "Backbone: patch embedding",
     "n_{patches} = ⌊(L − P) / S⌋ + 1 = ⌊(48 − 8) / 4⌋ + 1 = 11",
     "History L = 48 steps cut into patches of P = 8 with stride S = 4"),
    ("F3", "Backbone: Transformer encoder and NWP cross-attention",
     "Attention(Q, K, V) = softmax( Q K^T / √d_{k} ) V",
     "Lets PV patches look at the weather-forecast patches (8 heads, d_{model} = 128)"),
    ("F4", "Expert pool: expert disagreement",
     "σ_{expert} = √( mean_{h} Var_{k}( ŷ_{k,h} ) )",
     "Standard deviation across the 4 experts, averaged over the horizon (computed on detached outputs)"),
    ("F5", "Router: uncertainty fusion",
     "σ_{total} = α·n(σ_{NWP}) + (1 − α)·n(σ_{expert})\nα = sigmoid(w^T h_{expert} + b),  n(σ) = BatchNorm(log(1 + σ))",
     "Mixes the two uncertainty sources; α is learned from the embedding"),
    ("F6", "Router: gate with temperature and residual scale",
     "T = exp( clamp(softplus(β)·σ_{total}, ±8) )\nw = softmax( δ · z / T ),  δ(0) = 0",
     "More uncertainty → higher T → votes spread out. δ starts at 0 = uniform routing (ReZero idea)"),
    ("F7", "Router output with last-step shortcut",
     "ŷ_{routed} = Σ_{k} w_{k}·ŷ_{k} + ( W_{s} x_{last} + b_{s} )",
     "Weighted sum of the 4 experts plus a zero-initialised linear shortcut from the latest observation"),
]
FORMULAS2 = [
    ("F8", "Probabilistic decoder (Phase 5)",
     "P50 = clip( ŷ_{routed} + 0.1·tanh(g(h_{nwp}, σ_{total})), 0, 1 )\nlower_{h} = softplus( a_{h} + softplus(γ_{h})·σ_{total} )\nP10 = clip(P50 − lower),  P90 = clip(P50 + upper)",
     "Interval width grows with σ_{total}; P10 ≤ P50 ≤ P90 by construction"),
    ("F9", "Training: loss",
     "ρ_{τ}(e) = max( τ e, (τ − 1) e ),  e = y − ŷ\nL = L_{pinball} + 0.05 L_{sharp} + w(t) Σ_{k} L_{expert,k} + 0.01 L_{balance}",
     "Pinball loss for P10/P50/P90; sharpness; dense expert loss (w = 1 → 0.3); gate load balance"),
    ("F10", "Training: optimiser and weight averaging",
     "η_{t} = η_{min} + ½(η_{0} − η_{min})(1 + cos(π t / T))\nθ_{ema} ← d·θ_{ema} + (1 − d)·θ,  d = min(0.999, (1 + t)/(10 + t))",
     "Cosine learning-rate schedule (AdamW); EMA weights used for validation and checkpoints"),
    ("F11", "Evaluation",
     "nRMSE = 100·√mean((y − ŷ)²),   Skill = 1 − RMSE / RMSE_{persistence}\nPICP = mean(P10 ≤ y ≤ P90),  PINAW = mean(P90 − P10)\nWinkler = mean( width + (2/α)·(below + above) ),  α = 0.2",
     "Accuracy, skill vs 'same as now', coverage, width, and joint interval score"),
    ("F12", "Calibration (after training)",
     "E = max( P10 − y, y − P90 )  (daytime points, per regime)\nq̂ = Quantile_{⌈(n+1)(1−α)⌉/n}(E),   P10′ = P10 − q̂,  P90′ = P90 + q̂",
     "Conformalized quantile regression; brings daytime coverage to about 0.80"),
    ("F13", "Simulated Vellore site (demo only)",
     "T_{cell} = T_{amb} + 0.03125·GHI\nP = P_{peak}·(GHI/1000)·(1 − 0.004(T_{cell} − 25))·0.88·s,   P ≤ P_{peak}/1.25",
     "Temperature, soiling (s) and inverter limit; driven by ERA5 reanalysis, not the forecasts"),
    ("F14", "Comparison rule (results)",
     "“better” only if  | mean_{A} − mean_{B} | > 2·max(σ_{A}, σ_{B})",
     "A win is claimed only beyond about two run-to-run standard deviations"),
]

REFS_STUDIED = [
    "Y. Wang et al., “A distributed PV ultra-short-term power deterministic and probabilistic forecasting based on dynamic graph network with shape-amplitude criteria,” IEEE Trans. Ind. Appl., vol. 62, no. 2, pp. 1863–1873, 2026.",
    "H.-Y. Lo et al., “A novel QR-based probabilistic forecasting method for solar power generation,” IEEE Trans. Ind. Appl., vol. 61, no. 4, pp. 5381–5393, 2025.",
    "T. He, Y. Xu, and Y. Guo, “A probabilistic forecasting method for PV power under complex weather conditions based on multi-resolution deep learning,” in Proc. ICSGGE, 2026, pp. 125–130.",
    "H. Wang, R. Duan, and X. Yang, “A robust ensemble framework with context-gated dynamic fusion and temporal residual correction for PV power forecasting,” in Proc. ICEMCE, 2025, pp. 1022–1026.",
    "H. Zhao et al., “A two-stage ensemble learning framework for short-term nonparametric probabilistic PV power forecasting integrating multi-source NWPs,” IEEE Trans. Power Syst., vol. 41, no. 2, pp. 1231–1245, 2026.",
    "Y. He et al., “Day-ahead PV power probability density forecasting with GAN-QRCNLSTM based on high-resolution data reconstruction,” IEEE Trans. Ind. Informat., vol. 22, no. 2, pp. 1428–1439, 2026.",
    "M. Z. Heidrich et al., “Patch-based transformer with mixture of experts and conformal inference for curtailment time series forecasting,” IEEE Access, vol. 14, pp. 64144–64160, 2026.",
    "C. Lyu and S. Eftekharnejad, “Probabilistic solar generation forecasting for rapidly changing weather conditions,” IEEE Access, vol. 12, pp. 79091–79103, 2024.",
    "F. Lin et al., “Short-term probabilistic forecasting for regional PV power based on convolutional graph neural network and parameter transferring,” IEEE Trans. Power Syst., vol. 40, no. 3, pp. 2724–2736, 2025.",
    "C. Huang et al., “Spatial-temporal attention mechanism and graph convolutional networks based distributed PV ultra-short-term probabilistic power forecasting,” IEEE Trans. Ind. Appl., early access, 2026.",
    "R. Rastgoo et al., “A diffusion-based probabilistic ultra-short-term solar power prediction using the sky image sequences,” IEEE Access, vol. 13, pp. 149532–149553, 2025.",
    "L. O. Seman et al., “End-to-end decomposition-aware neural architecture search with sparse expert transformers for thermal energy forecasting,” IEEE Access, vol. 14, pp. 67765–67787, 2026.",
    "Q.-T. Phan, Y.-K. Wu, and Q.-D. Phan, “Enhancing one-day-ahead probabilistic solar power forecast with a hybrid transformer-LUBE model and missing data imputation,” IEEE Trans. Ind. Appl., vol. 60, no. 1, pp. 1396–1408, 2024.",
    "V. Jensen, F. M. Bianchi, and S. N. Anfinsen, “Ensemble conformalized quantile regression for probabilistic time series forecasting,” IEEE Trans. Neural Netw. Learn. Syst., vol. 35, no. 7, pp. 9014–9025, 2024.",
    "K. Song, M. Kim, and H. Kim, “Graph-based large scale probabilistic PV power forecasting insensitive to space-time missing data,” IEEE Trans. Sustain. Energy, vol. 16, no. 1, pp. 160–173, 2025.",
    "K. Wang et al., “Multimodal ensemble photovoltaic power forecasting incorporating diffusion-based satellite image prediction,” IEEE Trans. Sustain. Energy, vol. 17, no. 2, pp. 1425–1436, 2026.",
    "P. Zhao et al., “Probabilistic multienergy load forecasting based on hybrid attention-enabled transformer network and Gaussian process-aided residual learning,” IEEE Trans. Ind. Informat., vol. 20, no. 6, pp. 8379–8393, 2024.",
    "J. Ye et al., “TADNet: Temporal attention decomposition networks for probabilistic energy forecasting,” IEEE Trans. Power Syst., vol. 39, no. 6, pp. 7190–7202, 2024.",
    "W. Wang et al., “Ultra-short-term power probabilistic forecasting method for distributed photovoltaic based on joint distribution modeling,” IEEE Trans. Ind. Appl., vol. 62, no. 1, pp. 26–36, 2026.",
    "O. Abedinia et al., “Wind power forecasting enhancement utilizing adaptive quantile function and CNN-LSTM: A probabilistic approach,” IEEE Trans. Ind. Appl., vol. 60, no. 3, pp. 4446–4457, 2024.",
]
REFS_METHODS = [
    "A. Vaswani et al., “Attention is all you need,” in Proc. NeurIPS, 2017, pp. 5998–6008.   (F3)",
    "Y. Nie et al., “A time series is worth 64 words: Long-term forecasting with transformers,” in Proc. ICLR, 2023.   (F2)",
    "G. Ke et al., “LightGBM: A highly efficient gradient boosting decision tree,” in Proc. NeurIPS, 2017, pp. 3146–3154.   (baseline)",
    "R. Koenker and G. Bassett, “Regression quantiles,” Econometrica, vol. 46, no. 1, pp. 33–50, 1978.   (F9)",
    "R. L. Winkler, “A decision-theoretic approach to interval estimation,” J. Amer. Statist. Assoc., vol. 67, no. 337, pp. 187–191, 1972.   (F11)",
    "Y. Romano, E. Patterson, and E. J. Candès, “Conformalized quantile regression,” in Proc. NeurIPS, 2019, pp. 3543–3553.   (F12)",
    "T. Bachlechner et al., “ReZero is all you need: Fast convergence at large depth,” in Proc. UAI, PMLR 161, 2021, pp. 1352–1361.   (F6)",
    "B. T. Polyak and A. B. Juditsky, “Acceleration of stochastic approximation by averaging,” SIAM J. Control Optim., vol. 30, no. 4, pp. 838–855, 1992.   (F10)",
]


# ------------------------------------------------------------------ build
def main():
    prs = Presentation(str(TEMPLATE))
    S = list(prs.slides)
    (s_title, s_appr, s_aim, s_abs, s_lit, s_gap, s_obj, s_arch, s_func, s_mod, s_res, s_con, s_out, s_ref) = S
    foot = footers_from(s_aim)

    # ---------------- 1 title
    for sh in s_title.shapes:
        if sh.name == "Title 6":
            tf = sh.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            for txt, sz in (("B.Tech. – Computer Science and Engineering", 20), ("BCSE497J – Project-I", 20)):
                r = p.add_run(); r.text = txt; style_run(r, sz, True)
                p.add_line_break()
            r = p.add_run(); r.text = TITLE; style_run(r, 26, True)
        elif sh.name == "object 3":
            tf = sh.text_frame
            tf.clear()
            lines = [("Team members:", False, False), ("<<Name 1>>\t\t(<<Reg.No>>)", True, True), ("<<Name 2>>\t\t(<<Reg.No>>)", True, True),
                     ("<<Name 3>>\t\t(<<Reg.No>>)", True, True), ("", False, False), ("Faculty guide:", False, False), ("<<Guide name>>", True, True)]
            for i, (txt, b, hl) in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                r = p.add_run(); r.text = txt; style_run(r, 20, b, highlight=hl)

    # ---------------- 2 approval mail
    set_title(s_appr, "Approval Mail From Guide", 30)
    t = s_appr.shapes.title
    t.left, t.top, t.width, t.height = Cm(1.3), Cm(1.9), Cm(29.2), Cm(1.6)
    for sh in list(s_appr.shapes):
        if sh.shape_type == 13:       # sample picture from the template
            remove(sh)
    bp = body_ph(s_appr)
    place(bp, 1.5, 4.0, 30.5, 5.0)
    set_paragraphs(bp.text_frame, [
        ("Guide approval for the Review-3 points obtained by mail (screenshot below).", 0),
        ("Meeting with guide frequently: <<Satisfied>>   |   PowerPoint presentation: <<Approved>>   |   Final project report: <<Satisfied>>", 0, {"size": 16, "highlight": True}),
        ("Project demonstration: <<Approved>>   |   Project outcome: <<Ongoing>>   |   Turnitin plagiarism: <<__ %>>   |   AI score: <<__ %>> (should be ≤ 25%)", 0, {"size": 16, "highlight": True}),
    ], size=20)
    box = s_appr.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(4.0), Cm(9.6), Cm(25.8), Cm(7.4))
    box.fill.background()
    box.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    box.line.dash_style = 4
    tf = box.text_frame
    tf.clear()
    r = tf.paragraphs[0].add_run()
    r.text = "[ Paste the screenshot of the guide's approval mail here ]"
    style_run(r, 18, False, True, RGBColor(0x66, 0x66, 0x66), highlight=True)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # ---------------- 3 aim
    bp = body_ph(s_aim)
    place(bp, 2.3, 4.6, 29.2, 12.0)
    set_paragraphs(bp.text_frame, [
        "To design, build and evaluate ATERNET, a probabilistic solar-power forecaster that predicts the next 8 hours (30-minute steps) as a best guess (P50) with an 80% range (P10 to P90).",
        "To make it as accurate as the strongest baseline (LightGBM) while giving reliable, calibrated uncertainty ranges, and to test this honestly over several random runs and two regions.",
        "To demonstrate the system through a plain-language dashboard, including a clearly labelled simulated demonstration for VIT Vellore.",
    ], size=24, space_before=14)

    # ---------------- 4 abstract
    bp = body_ph(s_abs)
    place(bp, 2.3, 4.2, 29.2, 12.8)
    set_paragraphs(bp.text_frame, [
        "**Problem:** solar output depends on weather; grids need forecasts that also say how sure they are.".replace("**", ""),
        "ATERNET = PatchTST-style Transformer backbone + 4 weather-regime experts + uncertainty-aware router + probabilistic decoder (P10 ≤ P50 ≤ P90 always).",
        "Data: measured regional PV (Sheffield Solar PV-Live; Yorkshire and South West England, 2023–2026) + 5 NWP models (Open-Meteo). 74 automated tests pass.",
        f"Results: nRMSE {ms(Y[A], 'nrmse_pct')}% on Yorkshire (LightGBM {ms(Y['LightGBM quantile'], 'nrmse_pct')}, PatchTST {ms(Y['PatchTST quantile'], 'nrmse_pct')}); "
        f"{ms(SW[A], 'nrmse_pct')}% on South West (LightGBM {ms(SW['LightGBM quantile'], 'nrmse_pct')}).",
        f"Daytime coverage of the P10–P90 range: {SUMMARY['sites']['yorkshire']['calibration']['picp_daytime_before']:.2f} → {SUMMARY['sites']['yorkshire']['calibration']['picp_daytime_after']:.2f} (Yorkshire) and "
        f"{SUMMARY['sites']['south_west']['calibration']['picp_daytime_before']:.2f} → {SUMMARY['sites']['south_west']['calibration']['picp_daytime_after']:.2f} (South West) after calibration (target 0.80).",
        "Honest findings: the first learned gate hurt accuracy; only the dense expert loss measurably helped; the India demonstration is simulated.",
    ], size=19, space_before=8)

    # ---------------- 5 + 6 literature review (two slides, real tables)
    for sh in list(s_lit.shapes):
        if sh.shape_type == 13 or sh.name.startswith("Content Placeholder"):
            remove(sh)
    set_title(s_lit, "Literature Review (1/2)")
    s_lit.shapes.title.left, s_lit.shapes.title.top = Cm(1.3), Cm(1.0)
    add_table(s_lit, 1.2, 2.5, [4.3, 4.6, 7.6, 6.7, 8.5], LIT, size=8.5, row_h=1.2, header_size=11)
    s_lit2 = new_slide(prs, "Literature Review (2/2)", foot)
    s_lit2.shapes.title.top = Cm(1.0)
    add_table(s_lit2, 1.2, 2.5, [4.3, 4.6, 7.6, 6.7, 8.5], LIT2, size=8.5, row_h=1.2, header_size=11)

    # ---------------- 7 research gap
    bp = body_ph(s_gap)
    place(bp, 1.3, 4.0, 31.0, 13.2)
    set_paragraphs(bp.text_frame, [
        "Uncertainty from disagreement: none of the 20 papers uses the disagreement between several NWP models or between experts as an explicit uncertainty signal that also sets interval width and gate sharpness ([5] uses multi-source NWP as input, [4] gates on prediction variance, not both).",
        "Value of gating: mixtures of experts are reported for PV, curtailment and thermal energy ([4], [7], [12]), but whether a learned gate beats a plain average, and whether it is stable, needs a direct test for regional PV.",
        "Calibration where it matters: PV has exact zeros at night; coverage and conformal calibration should be done on daytime points.",
        "Data needs: graph and image methods ([1], [9], [10], [15], [11], [16]) need plant-level networks or image streams; a regional-aggregate + NWP model is useful where these are missing.",
        "Reliable comparison: differences of a few percent can be smaller than run-to-run variation; use several seeds and a second region.",
        "Regions without open data (e.g. our campus): a clearly labelled simulated demonstration is needed.",
    ], size=17, space_before=8)

    # ---------------- 8 objectives
    bp = body_ph(s_obj)
    place(bp, 1.3, 3.7, 31.0, 8.6)
    set_paragraphs(bp.text_frame, [
        "O1  Route between weather-regime experts using the combined uncertainty (context-gated routing)",
        "O2  Make the interval width depend on uncertainty; P10 ≤ P50 ≤ P90 by construction",
        "O3  Train all parts jointly; keep every expert learning (dense expert loss)",
        "O4  Evaluate per weather regime and per forecast horizon",
        "O5  Compare honestly with LightGBM, PatchTST, LSTM over 3 seeds and on a second region",
        "O6  Deliver a tested, reproducible system, a plain-language dashboard and a simulated Vellore demo",
    ], size=17, space_before=5)
    for sh in s_obj.shapes:
        if sh.name == "TextBox 6":
            tf = sh.text_frame
            tf.clear()
            lines = ["Identified SDGs: SDG 7 – Affordable and Clean Energy;  SDG 13 – Climate Action",
                     "Technology Readiness Level: TRL 4 (validated in a research environment on real data)",
                     "Outcomes: project report + dashboard; conference paper planned"]
            for i, txt in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                r = p.add_run(); r.text = txt; style_run(r, 14, True)
            sh.left, sh.top, sh.width, sh.height = Cm(1.3), Cm(10.2), Cm(31.0), Cm(2.6)
        if sh.name == "Picture 7":
            sh.top = Cm(13.4); sh.height = Cm(4.2); sh.width = Cm(5.3)
        if sh.name == "Picture 8":
            sh.top = Cm(13.4); sh.height = Cm(4.2); sh.width = Cm(4.6)

    # ---------------- 9 architecture (two diagrams)
    set_title(s_arch, "Framework / Architecture / Block Diagram")
    bp = body_ph(s_arch)
    remove(bp)
    add_picture_fit(s_arch, FIG / "fig_4_1_system_architecture.png", 1.0, 3.7, 15.0, 13.6)
    add_picture_fit(s_arch, FIG / "fig_4_2_model_architecture.png", 16.6, 3.7, 16.3, 13.6)
    add_text(s_arch, 1.0, 17.2, 15.0, 0.8, [("System architecture: data → model/experiment → application layers", 0)], size=11, align=PP_ALIGN.CENTER)
    add_text(s_arch, 16.6, 17.2, 16.3, 0.8, [("ATERNET model: backbone → experts → router → decoder", 0)], size=11, align=PP_ALIGN.CENTER)

    # ---------------- 10 formula positions (native diagram)
    s_pos = new_slide(prs, "Where Each Formula Is Used (F1 – F14)", foot)
    blocks = [  # (label, formulas, x, y, w, fill)
        ("Data\npreparation\n(Phase 1)", "F1", 0.9, 4.6, 4.1, "DCE9F7"),
        ("Backbone\n(Phase 2)", "F2, F3", 5.6, 4.6, 4.1, "DCE9F7"),
        ("Expert pool\n(Phase 3)", "F4", 10.3, 4.6, 4.1, "DDEFE0"),
        ("Adaptive\nrouter\n(Phase 4)", "F5, F6, F7", 15.0, 4.6, 4.1, "FBEBC8"),
        ("Probabilistic\ndecoder\n(Phase 5)", "F8", 19.7, 4.6, 4.1, "F8DCDC"),
        ("Output\nP10, P50, P90", "", 24.4, 4.6, 3.6, "FFFFFF"),
    ]
    for label, tag, x, y, w, fill in blocks:
        b = s_pos.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(3.2))
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        b.line.color.rgb = BLUE
        tf = b.text_frame; tf.clear(); tf.word_wrap = True
        for i, line in enumerate(label.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line; style_run(r, 13, i == 0, color=RGBColor(0, 0, 0))
        if tag:
            tg = s_pos.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y + 3.35), Cm(w), Cm(0.8))
            tg.fill.solid(); tg.fill.fore_color.rgb = BLUE; tg.line.fill.background()
            tf = tg.text_frame; tf.clear()
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run(); r.text = tag; style_run(r, 13, True, color=RGBColor(255, 255, 255))
    for xa in (5.0, 9.7, 14.4, 19.1, 23.8):
        c = s_pos.shapes.add_connector(1, Cm(xa), Cm(6.2), Cm(xa + 0.6), Cm(6.2))
        c.line.color.rgb = RGBColor(0x33, 0x33, 0x33)
        c.line.width = Pt(1.5)
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle")
    # second row: what happens around the model
    row2 = [
        ("Training\n(loss, optimiser, EMA)", "F9, F10", 0.9, 10.2, 7.6, "EEEEEE"),
        ("Evaluation\n(metrics, comparison rule)", "F11, F14", 9.2, 10.2, 7.6, "EEEEEE"),
        ("Calibration\n(daytime conformal)", "F12", 17.5, 10.2, 7.6, "EEEEEE"),
        ("Simulated Vellore site\n(demo data only)", "F13", 25.8, 10.2, 6.9, "EEEEEE"),
    ]
    for label, tag, x, y, w, fill in row2:
        b = s_pos.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(2.6))
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        b.line.color.rgb = RGBColor(0x55, 0x55, 0x55)
        tf = b.text_frame; tf.clear(); tf.word_wrap = True
        for i, line in enumerate(label.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line; style_run(r, 13, i == 0, color=RGBColor(0, 0, 0))
        tg = s_pos.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y + 2.75), Cm(w), Cm(0.8))
        tg.fill.solid(); tg.fill.fore_color.rgb = BLUE; tg.line.fill.background()
        tf = tg.text_frame; tf.clear(); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run(); r.text = tag; style_run(r, 13, True, color=RGBColor(255, 255, 255))
    add_text(s_pos, 0.9, 3.6, 27.0, 0.9, [("Model forward pass (one forecast)", 0, {"bold": True})], size=14)
    add_text(s_pos, 0.9, 9.3, 27.0, 0.9, [("Around the model", 0, {"bold": True})], size=14)
    add_text(s_pos, 0.9, 14.9, 32.0, 2.6, [
        ("Blue tags = formula numbers; the next two slides give each formula and what it does.", 0),
        ("Every formula here is implemented in the project code (src/phase1–5, src/training, src/evaluation); F2/F3 follow PatchTST and the Transformer, F6 follows ReZero, F12 follows CQR.", 0),
    ], size=13, bullets=True)

    # ---------------- 11 + 12 formula tables
    def formula_slide(title, items, notes):
        s = new_slide(prs, title, foot)
        s.shapes.title.top = Cm(1.0)
        rows_ = [["No.", "Position in the system", "Formula", "What it does"]] + [[a, b, c, d] for a, b, c, d in items]
        add_table(s, 0.9, 2.7, [1.4, 6.2, 14.5, 9.6], rows_, size=11, rich_cols=(2, 3), header_size=13, bold_first_col=True)
        s.notes_slide.notes_text_frame.text = notes
        return s
    formula_slide("Formulas and Positions (1/2): Model Forward Pass", FORMULAS,
                  "Walk through the forward pass left to right: F1 prepares the inputs, F2-F3 are in the backbone, F4 measures expert disagreement, F5-F7 are in the router, and F8 (next slide) is the decoder.")
    formula_slide("Formulas and Positions (2/2): Decoder, Training, Evaluation", FORMULAS2,
                  "F8 is the decoder; F9-F10 are used during training; F11 and F14 in evaluation and comparison; F12 after training for calibration; F13 only for the simulated Vellore demonstration.")

    # ---------------- 13 functional requirements
    bp = body_ph(s_func)
    place(bp, 1.3, 4.0, 31.0, 13.2)
    set_paragraphs(bp.text_frame, [
        "Download measured regional PV (PV-Live) and five-model weather forecasts (Open-Meteo); align to a 30-minute grid",
        "Clean, derive features (clear-sky index, time encodings, forecast spread, regime labels) and build 24 h → 8 h windows without look-ahead",
        "Train ATERNET jointly (early stopping, checkpoints, weight averaging); output P10, P50, P90 for 16 steps with P10 ≤ P50 ≤ P90",
        "Evaluate: nRMSE, nMAE, R², PICP, PINAW, Winkler, skill vs persistence; per regime and per horizon",
        "Train and evaluate LightGBM, PatchTST and LSTM baselines under the same protocol",
        "Calibrate the intervals per regime on daytime validation points",
        "Run multi-seed experiments and ablations; aggregate mean ± standard deviation",
        "Live forecast and backtest commands; interactive dashboard; second region and simulated site by configuration only",
    ], size=17, space_before=8)

    # ---------------- 14 modules + dashboard screenshot
    bp = body_ph(s_mod)
    place(bp, 1.0, 3.9, 19.6, 13.4)
    set_paragraphs(bp.text_frame, [
        "1  Data pipeline (Phase 1): PV-Live + Open-Meteo → clean → features → windows (48 → 16) → chronological 70/15/15 split",
        "2  Temporal backbone (Phase 2): patches, 3-layer Transformer, NWP cross-attention, σ encoder",
        "3  Expert pool (Phase 3): dawn-ramp, clear-peak, overcast, volatile experts + σ_expert",
        "4  Adaptive router (Phase 4): uncertainty fusion, residual gate, last-step shortcut",
        "5  Probabilistic decoder (Phase 5): P50 + widths → P10, P90",
        "6  Training and calibration: composite loss, AdamW, EMA, daytime conformal calibration",
        "7  Baselines and evaluation: LightGBM, PatchTST, LSTM; metrics; multi-seed runner",
        "8  Dashboard and live forecast: Streamlit (9 pages), predict.py; 74 automated tests",
    ], size=14, space_before=6)
    add_picture_fit(s_mod, FIG / "crop_better_1.png", 21.0, 3.9, 12.0, 13.4)
    add_text(s_mod, 21.0, 17.2, 12.0, 0.7, [("Dashboard: “Is ATERNET better?” page", 0)], size=11, align=PP_ALIGN.CENTER)

    # ---------------- 15 + 16 results
    set_title(s_res, "Experiments and Results (1/2): Comparison")
    bp = body_ph(s_res)
    remove(bp)
    add_picture_fit(s_res, FIG / "fig_7_1_comparison_yorkshire.png", 0.8, 3.7, 16.2, 8.0)
    add_picture_fit(s_res, FIG / "fig_7_2_comparison_south_west.png", 17.0, 3.7, 16.2, 8.0)
    add_text(s_res, 1.0, 11.9, 32.0, 5.6, [
        ("Setup: test period 8 Mar – 29 Sep 2026; 3 training runs per model (LSTM: 6); mean ± std; a difference counts only beyond 2 std (F14).", 0),
        (f"Yorkshire: ATERNET {ms(Y[A], 'nrmse_pct')}% = LightGBM {ms(Y['LightGBM quantile'], 'nrmse_pct')}% (tie), better than PatchTST {ms(Y['PatchTST quantile'], 'nrmse_pct')}%; Winkler {ms(Y[A], 'winkler', 4)} vs {ms(Y['LightGBM quantile'], 'winkler', 4)} / {ms(Y['PatchTST quantile'], 'winkler', 4)}.", 0),
        (f"South West (no re-tuning): ATERNET {ms(SW[A], 'nrmse_pct')}% vs LightGBM {ms(SW['LightGBM quantile'], 'nrmse_pct')}% and PatchTST {ms(SW['PatchTST quantile'], 'nrmse_pct')}%.", 0),
        ("LSTM is unstable: 3 of 6 runs failed (nRMSE 7.0–9.2%); ATERNET varies by only 0.02.", 0),
    ], size=15, bullets=True)
    s_res2 = new_slide(prs, "Experiments and Results (2/2): What Helped and Calibration", foot)
    s_res2.shapes.title.top = Cm(1.4)
    add_picture_fit(s_res2, FIG / "fig_7_3_model_development.png", 0.5, 3.4, 11.4, 6.6)
    add_picture_fit(s_res2, FIG / "fig_7_4_ablation.png", 11.9, 3.4, 11.4, 6.6)
    add_picture_fit(s_res2, FIG / "fig_7_5_daytime_coverage.png", 23.3, 3.4, 10.2, 6.6)
    add_text(s_res2, 0.8, 10.4, 32.2, 7.2, [
        ("Development (Yorkshire): original learned gate 4.39 ± 0.20 → residual gate 4.12 → + last-step shortcut 3.83 → + EMA 3.70 ± 0.02 nRMSE (LightGBM 3.72).", 0),
        ("Ablations (switch one part off): only the dense expert loss clearly mattered (Winkler +8%, coverage 0.82 → 0.78); gate, temperature, fixed width, single expert, regime weights made no measurable difference.", 0),
        ("Calibration (F12): daytime coverage 0.67 → 0.79 (Yorkshire) and 0.73 → 0.80 (South West); simulated Vellore only 0.70 → 0.74.", 0),
        ("Backtest 1–16 Oct 2026: daylight MAE 42 MW vs 147 MW for day-ahead persistence (3.5× better). Vellore results are simulated (single run), not real-world accuracy.", 0),
    ], size=14, bullets=True)
    s_res2.notes_slide.notes_text_frame.text = "Left: nRMSE at each development stage. Middle: effect of removing each model part (red = clearly worse). Right: daytime coverage before and after calibration against the 80% target."

    # ---------------- 17 conclusion
    bp = body_ph(s_con)
    place(bp, 1.3, 3.9, 31.0, 13.3)
    set_paragraphs(bp.text_frame, [
        ("Conclusion", 0, {"bold": True, "nobullet": True, "size": 20}),
        ("ATERNET matches LightGBM on Yorkshire (3.70 vs 3.72 nRMSE), beats it on South West England (2.63 vs 2.69) and beats PatchTST on both; its ranges score better on Winkler.", 1),
        ("A residual gate, a last-step shortcut and EMA fixed the harm done by the first learned gate; training is stable (std 0.02).", 1),
        ("Daytime-only calibration brings coverage to about 0.80; all-hours coverage had hidden under-coverage.", 1),
        ("Adaptive routing and regime specialisation did not improve accuracy; only the dense expert loss did. The India results are simulated.", 1),
        ("Future work", 0, {"bold": True, "nobullet": True, "size": 20, "before": 14}),
        ("Obtain measured PV data for an Indian site (e.g. a campus plant) and fine-tune / retrain.", 1),
        ("More regions and seasons; tune the baselines equally; ensembles of several ATERNET runs.", 1),
        ("Make the experts specialise genuinely (diversity penalty, regime-specific data) and re-test routing; extend the horizon and deploy as a service.", 1),
    ], size=17, level_sizes={1: 16}, space_before=6)

    # ---------------- 18 outcome achieved
    set_title(s_out, "Outcome Achieved")
    bp = body_ph(s_out)
    set_paragraphs(bp.text_frame, [
        ("Patent: <<not filed>>   |   Conference (Scopus): <<planned – paper on ATERNET results>>   |   Journal: <<not submitted>>", 0, {"size": 16, "highlight": True}),
        ("Start-up / product: <<none>>", 0, {"size": 16, "highlight": True}),
        ("Others: project report (71 pages), working Streamlit dashboard, 74 automated tests, reproducible build scripts", 0, {"size": 16}),
        ("Update this slide with proof (screenshot) of any accepted / published outcome before uploading.", 0, {"size": 16, "highlight": True}),
    ], size=16, space_before=10)
    place(bp, 2.3, 5.0, 29.2, 8.0)
    # keep the template's red upload reminder, which sits in the same placeholder in the original; add it back as a text box
    add_text(s_out, 2.3, 13.4, 29.2, 2.8, [
        ("Outcome Achieved proof to be uploaded on or before 27 Oct 2026 in the given link: https://forms.gle/B9B71VPy5b9bZ9dJ8", 0, {"color": RED, "size": 15, "nobullet": True}),
    ], size=15)

    # ---------------- 19 references
    set_title(s_ref, "References")
    bp = body_ph(s_ref)
    remove(bp)
    half = len(REFS_STUDIED) // 2
    def col(items, start):
        return [(f"[{start + i}] {t}", 0, {"nobullet": True, "before": 2}) for i, t in enumerate(items)]
    add_text(s_ref, 0.8, 3.5, 16.2, 14.0, col(REFS_STUDIED[:half], 1), size=8.5)
    methods = col(REFS_METHODS, 21)
    add_text(s_ref, 17.0, 3.5, 16.2, 14.0, col(REFS_STUDIED[half:], half + 1) + [("Methods used in the formulas:", 0, {"bold": True, "nobullet": True, "before": 6})] + methods, size=8.5)

    # ---------------- final order (19 slides)
    order = [s_title, s_appr, s_aim, s_abs, s_lit, s_lit2, s_gap, s_obj, s_arch, s_pos]
    # formula slides were appended after s_pos; collect by title
    by_title = {sl.shapes.title.text_frame.text: sl for sl in prs.slides if sl.shapes.title is not None}
    order += [by_title["Formulas and Positions (1/2): Model Forward Pass"],
              by_title["Formulas and Positions (2/2): Decoder, Training, Evaluation"],
              s_func, s_mod, s_res, s_res2, s_con, s_out, s_ref]
    assert len(order) == len(prs.slides) == 19, (len(order), len(prs.slides))
    lst = prs.slides._sldIdLst
    id_by_slide = {}
    for sldId in list(lst):
        id_by_slide[prs.part.related_part(sldId.rId)] = sldId
    for sldId in list(lst):
        lst.remove(sldId)
    for sl in order:
        lst.append(id_by_slide[sl.part])
    prs.save(str(OUT))
    print(f"saved {OUT.name}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
